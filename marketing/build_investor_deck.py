from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = "终焉_投资人介绍.pptx"
W, H = 13.333, 7.5
BG = RGBColor(7, 6, 11)
PAPER = RGBColor(242, 234, 216)
MUTED = RGBColor(168, 159, 141)
GOLD = RGBColor(227, 194, 124)
BLUE = RGBColor(139, 147, 248)
GREEN = RGBColor(78, 203, 156)
RED = RGBColor(238, 106, 114)
AMBER = RGBColor(242, 169, 59)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape

def textbox(slide, text, x, y, w, h, size=20, color=PAPER, bold=False, font="Microsoft YaHei", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    f = run.font; f.name = font; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    return box

def base(title, kicker):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0.55, 0.42, 0.08, 0.52, GOLD)
    textbox(slide, kicker.upper(), 0.82, 0.39, 8, 0.24, 10, GOLD, True, "Consolas")
    textbox(slide, title, 0.82, 0.72, 11.8, 0.55, 27, PAPER, True)
    textbox(slide, f"终焉  /  TEN DAYS GAMBIT", 10.4, 0.45, 2.4, 0.24, 9, MUTED, False, "Consolas", PP_ALIGN.RIGHT)
    return slide

def bullet(slide, text, x, y, w, color=PAPER, size=18, accent=GOLD):
    rect(slide, x, y+0.13, 0.08, 0.08, accent)
    textbox(slide, text, x+0.22, y, w-0.22, 0.45, size, color)

def footer(slide, n):
    textbox(slide, f"{n:02d}", 12.25, 7.05, 0.5, 0.2, 10, MUTED, False, "Consolas", PP_ALIGN.RIGHT)

# 1
s = prs.slides.add_slide(prs.slide_layouts[6]); rect(s,0,0,W,H,BG)
textbox(s, "终焉", 0.9, 1.1, 7, 1.1, 62, PAPER, True)
textbox(s, "可接入的分布式智能游戏世界", 0.95, 2.25, 8.6, 0.6, 28, GOLD, True)
textbox(s, "真人、Agent 与裁判 Agent，在同一套规则里共同上桌。", 0.98, 3.1, 8.8, 0.45, 21, MUTED)
for i,(label,c) in enumerate([("人类",BLUE),("Agent",GREEN),("规则",AMBER),("裁判",RED)]):
    x=8.6+(i%2)*1.65; y=1.3+(i//2)*1.65
    rect(s,x,y,1.32,1.32,BG,c,True); textbox(s,label,x,y+0.48,1.32,0.25,18,c,True,align=PP_ALIGN.CENTER)
rect(s,0.98,5.75,4.0,0.55,RGBColor(25,31,33),RGBColor(95,75,42)); textbox(s,"黑客松首发：可现场游玩、可观战、可接入",1.22,5.91,3.55,0.18,13,GOLD,True)
textbox(s,"github.com/icelikey/shirizhongyan",0.98,6.75,5,0.22,12,MUTED,False,"Consolas"); footer(s,1)

# 2
s=base("AI 游戏缺少一张“同桌”", "opportunity")
for i,(head,body,c) in enumerate([
    ("只会聊天", "策略停留在文本里，难以进入共享状态和真实对局。", BLUE),
    ("看不懂结果", "观众看到输出，却看不到决策如何改变局面。", RED),
    ("无法验证", "缺少可回放的规则、裁判与事件证据。", AMBER)]):
    x=0.9+i*4.05; rect(s,x,1.75,3.55,2.2,RGBColor(19,25,27),c,True); textbox(s,head,x+0.3,2.08,2.9,0.35,22,c,True); textbox(s,body,x+0.3,2.7,2.9,0.8,17,PAPER); rect(s,x+0.3,3.72,2.4,0.06,c)
textbox(s,"终焉把 AI 从“对话角色”变成“承担决策与结果的同桌玩家”。",0.95,5.1,11.2,0.6,27,GOLD,True)
footer(s,2)

# 3
s=base("终焉：让策略被看见", "product")
bullet(s,"人类可以亲自玩，也可以带着自己的 Agent 玩。",1.0,1.7,5.5,accent=BLUE)
bullet(s,"外部 Agent 通过公开 CLI / HTTP 协议注册、入座、观测、行动。",1.0,2.35,6.1,accent=GREEN)
bullet(s,"算法处理确定性规则，Jev / AI / Agent 组成 3/5/7 奇数裁判团。",1.0,3.0,6.3,accent=RED)
bullet(s,"每局产生事件、结算、奖励与高光，观众能追踪因果链。",1.0,3.65,6.0,accent=AMBER)
rect(s,8.0,1.75,3.8,3.25,RGBColor(19,25,27),GOLD,True)
for i,t in enumerate(["注册 Agent","进入房间","有限视角","策略行动","裁决回放"]):
    y=2.05+i*0.55; textbox(s,t,8.35,y,2.9,0.25,17,PAPER,True,align=PP_ALIGN.CENTER)
    if i<4: textbox(s,"↓",9.68,y+0.25,0.25,0.2,14,GOLD,True,align=PP_ALIGN.CENTER)
footer(s,3)

# 4
s=base("一条可重复的 AI 游戏闭环", "loop")
steps=[("01","入座",BLUE),("02","观察",GREEN),("03","决策",AMBER),("04","裁决",RED),("05","进化",GOLD)]
for i,(num,label,c) in enumerate(steps):
    x=0.95+i*2.45; rect(s,x,2.15,1.45,1.45,BG,c,True); textbox(s,num,x,2.43,1.45,0.25,13,c,False,"Consolas",PP_ALIGN.CENTER); textbox(s,label,x,2.85,1.45,0.35,22,PAPER,True,align=PP_ALIGN.CENTER)
    if i<4: textbox(s,"→",x+1.55,2.67,0.65,0.3,22,MUTED,True,align=PP_ALIGN.CENTER)
textbox(s,"上下文、动作、规则和结果都留下结构化证据。",1.0,4.55,9.8,0.45,24,GOLD,True)
textbox(s,"这使“意外行为”能够被复现、被观众理解，并成为下一局策略的输入。",1.0,5.2,10.8,0.4,18,MUTED)
footer(s,4)

# 5
s=base("首发内容：少而有辨识度", "game packs")
rows=[("青野算庭 · 猜平均数","已服务端联机","level-k 思考 / 结果清晰",GREEN),("红眼病 · 少数派票决","已服务端联机","群体心理 / 反直觉",GREEN),("规则辩论","下一阶段","奇数裁判 / 规则质询",RED),("超能力赛马","下一阶段","技能卡 / 赛道事件 / 3D 高光",AMBER),("月影狼人杀","本地可玩原型","后续服务端化 / 语音",BLUE)]
for i,(a,b,c,col) in enumerate(rows):
    y=1.55+i*0.85; rect(s,0.95,y,2.95,0.58,RGBColor(19,25,27),col,True); textbox(s,a,1.18,y+0.17,2.55,0.22,15,PAPER,True); textbox(s,b,4.35,y+0.17,2.1,0.22,14,col,True); textbox(s,c,7.05,y+0.17,4.4,0.22,15,MUTED)
textbox(s,"每个内容包复用同一套 Gateway、状态、裁判、奖励与回放基础设施。",0.98,6.05,11.1,0.35,20,GOLD,True)
footer(s,5)

# 6
s=base("技术壁垒在“可接入 + 可验证”", "moat")
cards=[("TDG-WP", "任何语言、任何框架的 Agent 都能上桌。", BLUE), ("有限视角", "服务端保管真实状态，每个席位只看到被允许看到的内容。", GREEN), ("奇数裁判", "算法、Jev、Agent 各司其职，语义争议有复核路径。", RED), ("事件回放", "把策略输入、规则效果与高光画面串成一条证据链。", AMBER)]
for i,(h,b,c) in enumerate(cards):
    x=0.95+(i%2)*6.0; y=1.65+(i//2)*2.25; rect(s,x,y,5.4,1.55,RGBColor(19,25,27),c,True); textbox(s,h,x+0.3,y+0.25,1.9,0.28,20,c,True); textbox(s,b,x+0.3,y+0.75,4.7,0.5,16,PAPER)
footer(s,6)

# 7
s=base("商业化：从活动房间到内容生态", "business")
for i,(h,b,c) in enumerate([
    ("事件与赛事", "线下黑客松、品牌活动、Agent 对抗赛", BLUE),
    ("内容包", "规则、事件、能力、视觉资产组成可发布玩法", GREEN),
    ("创作者工具", "让设计者发布游戏，平台提供协议、裁判和回放", AMBER),
    ("教育与训练", "把谈判、判断、协作和 AI 调教变成可观察体验", RED)]):
    x=0.95+i*3.0; rect(s,x,1.8,2.55,2.8,RGBColor(19,25,27),c,True); textbox(s,h,x+0.22,2.15,2.1,0.35,19,c,True); textbox(s,b,x+0.22,2.9,2.1,1.0,16,PAPER)
textbox(s,"同一套基础设施支持更多游戏、更多 Agent 和更多观众。",0.98,5.55,10.8,0.4,23,GOLD,True)
footer(s,7)

# 8
s=base("黑客松现场：先证明体验，再证明规模", "demo")
targets=[("1 局", "评委从网页进入并完成对局"),("5 分钟", "外部 Agent 完成注册、入座、行动"),("1 条链", "观众看懂输入 → 决策 → 裁决 → 结果")]
for i,(a,b) in enumerate(targets):
    x=1.0+i*4.0; rect(s,x,1.85,3.2,2.15,RGBColor(19,25,27),GOLD,True); textbox(s,a,x,2.25,3.2,0.6,34,GOLD,True,align=PP_ALIGN.CENTER); textbox(s,b,x+0.28,3.2,2.64,0.5,16,PAPER,False,align=PP_ALIGN.CENTER)
bullet(s,"现场可玩：两款确定性联机游戏已经提供真实闭环。",1.0,5.05,10.5,accent=GREEN)
bullet(s,"现场可讲：规则辩论与超能力赛马展示下一阶段的 AI 深度和视觉高光。",1.0,5.62,11.0,accent=AMBER)
footer(s,8)

# 9
s=base("路线图：可靠性先于扩张", "roadmap")
road=[("现在","两款联机游戏 / Agent CLI / Discovery",GREEN),("下一阶段","规则辩论 / 赛马 / 事件真实落库 / 观战回放",AMBER),("产品化","狼人杀服务端化 / 语音 / 3D / 能力进化",BLUE),("规模化","幂等、outbox、租约、状态服务、多实例",RED)]
for i,(h,b,c) in enumerate(road):
    y=1.55+i*1.1; rect(s,1.0,y,1.6,0.62,c,c,True); textbox(s,h,1.0,y+0.18,1.6,0.2,15,BG,True,align=PP_ALIGN.CENTER); textbox(s,b,3.05,y+0.16,8.8,0.25,18,PAPER)
footer(s,9)

# 10
s=base("我们要构建的不是一局游戏", "closing")
textbox(s,"而是一套让数字生命共同进入世界的基础设施。",1.0,1.65,11.4,0.6,30,GOLD,True)
textbox(s,"终焉把“AI 会做什么”变成观众可以参与、验证和记住的事件。",1.0,2.7,10.8,0.45,22,PAPER)
for i,t in enumerate(["可接入", "可观战", "可裁决", "可进化"]):
    x=1.0+i*2.85; rect(s,x,4.1,2.25,0.9,RGBColor(19,25,27),[BLUE,GREEN,RED,AMBER][i],True); textbox(s,t,x,4.39,2.25,0.25,19,[BLUE,GREEN,RED,AMBER][i],True,align=PP_ALIGN.CENTER)
textbox(s,"寻求：赛事合作 · 内容共创 · Agent 生态伙伴 · 产品化资源",1.0,6.25,10.5,0.35,18,MUTED)
footer(s,10)

prs.save(OUT)
print(OUT)
